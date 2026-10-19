"""Testa o servidor multi-cliente contra uma Graph API falsa."""
import asyncio, json, os, random, sys, tempfile
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import parse_qs, urlparse
import httpx

sys.path.insert(0, sys.argv[1])
for k in ("META_ACCESS_TOKEN", "META_AD_ACCOUNT_ID", "META_EXPORT_DIR"):
    os.environ.pop(k, None)
import server  # noqa: E402

tmp = Path(tempfile.mkdtemp())
server.BASE_DIR, server.ENV_FILE, server.CLIENTES_DIR = tmp, tmp / ".env", tmp / "clientes"
server._BACKOFF = (0, 0, 0); server.PAGINA = 40
(tmp / "clientes").mkdir()
TOK = {"primos": "TOKEN_PRIMOS_aaa111", "donaco": "TOKEN_DONACO_bbb222", "ruim": "TOKEN_RUIM_ccc333"}
ACC = {"primos": "111111111111111", "donaco": "333333333333333"}
(tmp / "clientes/primos.env").write_text(f"NOME=Primos Premiações\nMETA_ACCESS_TOKEN={TOK['primos']}\nMETA_AD_ACCOUNT_ID={ACC['primos']}\n")
(tmp / "clientes/donaco.env").write_text(f"NOME=Dona Co\nMETA_ACCESS_TOKEN={TOK['donaco']}\nMETA_AD_ACCOUNT_ID=\n")
DONO = {v: k for k, v in TOK.items() if k != "ruim"}

random.seed(7)
D1, D2 = date(2026, 9, 1), date(2026, 9, 14)
def ads(prefix, n):
    return [(f"{prefix}0000000000000{i}1", f"[VENDAS] {prefix} camp {i}", "OUTCOME_SALES", f"{prefix}0000000000000{i}2", f"conj {i}",
             f"{prefix}0000000000000{i}3", f"video_{i}", "ARCHIVED" if i == 2 else "ACTIVE") for i in range(1, n + 1)]
ADS = {"primos": ads("1202", 3), "donaco": ads("1303", 2)}
AGES, GENDERS, PLATS = ["18-24", "25-34", "35-44"], ["female", "male"], ["facebook", "instagram"]
FACTS = {}
for cli, lst in ADS.items():
    FACTS[cli] = []
    d = D1
    while d <= D2:
        for ad in lst:
            for a in AGES:
                for g in GENDERS:
                    for p in PLATS:
                        imp = random.randint(200, 900); buys = random.randint(0, 4)
                        FACTS[cli].append(dict(ad=ad, day=d, age=a, gender=g, plat=p, spend=round(random.uniform(5, 40), 2), imp=imp,
                                               people=set(random.sample(range(60000), k=int(imp * .8))), cl=random.randint(3, 30),
                                               buys=buys, val=buys * random.choice([20, 25, 40])))
        d += timedelta(days=1)
LOG = []  # (dono do token, caminho)
STATE = {"rate": False, "too_much": 0}

def err(status, code, msg): return httpx.Response(status, json={"error": {"message": msg, "code": code}})

def agg(facts, level, daily, bks, since, until, acc):
    groups = {}
    for f in facts:
        key = []
        if level in ("campaign", "adset", "ad"): key.append(f["ad"][0])
        if level in ("adset", "ad"): key.append(f["ad"][3])
        if level == "ad": key.append(f["ad"][5])
        if daily: key.append(f["day"])
        for b in bks: key.append({"age": f["age"], "gender": f["gender"], "publisher_platform": f["plat"]}[b])
        groups.setdefault(tuple(key), []).append(f)
    rows = []
    for key in sorted(groups, key=str):
        fs = groups[key]; ad = fs[0]["ad"]; imp = sum(x["imp"] for x in fs); reach = len(set().union(*(x["people"] for x in fs)))
        buys = sum(x["buys"] for x in fs); val = sum(x["val"] for x in fs)
        row = {"account_id": acc, "spend": f"{sum(x['spend'] for x in fs):.2f}", "impressions": str(imp), "reach": str(reach),
               "frequency": f"{imp / reach:.6f}", "clicks": "9", "inline_link_clicks": str(sum(x["cl"] for x in fs)),
               "date_start": (fs[0]["day"] if daily else since).isoformat(), "date_stop": (fs[0]["day"] if daily else until).isoformat()}
        if buys:
            row["actions"] = [{"action_type": t, "value": str(buys)} for t in ("purchase", "omni_purchase", "offsite_conversion.fb_pixel_purchase")]
            row["action_values"] = [{"action_type": t, "value": f"{val:.2f}"} for t in ("purchase", "omni_purchase")]
        if level in ("campaign", "adset", "ad"): row.update(campaign_id=ad[0], campaign_name=ad[1], objective=ad[2])
        if level in ("adset", "ad"): row.update(adset_id=ad[3], adset_name=ad[4])
        if level == "ad": row.update(ad_id=ad[5], ad_name=ad[6])
        for b in bks: row[b] = {"age": fs[0]["age"], "gender": fs[0]["gender"], "publisher_platform": fs[0]["plat"]}[b]
        rows.append(row)
    return rows

def handler(req):
    assert req.method == "GET", req.method
    tok = req.headers.get("authorization", "").removeprefix("Bearer ")
    u = urlparse(str(req.url)); q = parse_qs(u.query)
    assert not any(t in str(req.url) for t in TOK.values()), "token em URL"
    assert "NAO_SEGUIR" not in str(req.url)
    cli = DONO.get(tok); p = u.path.removeprefix("/v26.0/")
    LOG.append((cli, p))
    if cli is None: return err(400, 190, "Invalid OAuth access token - Cannot parse access token")
    acc = ACC[cli]
    if p == "me": return httpx.Response(200, json={"id": "9", "name": f"sysuser-{cli}"})
    if p == "me/permissions": return httpx.Response(200, json={"data": [{"permission": "ads_read", "status": "granted"}]})
    if p == "me/adaccounts":
        return httpx.Response(200, json={"data": [{"account_id": acc, "name": f"{cli.upper()} BM01 CA 01", "account_status": 1, "currency": "BRL", "timezone_name": "America/Bahia"}]})
    if p.startswith("act_") and not p.startswith(f"act_{acc}"):
        return err(403, 200, "(#200) Ad account owner has NOT grant ads_management or ads_read permission")
    if p == f"act_{acc}": return httpx.Response(200, json={"account_id": acc, "name": f"{cli.upper()} BM01 CA 01", "currency": "BRL", "timezone_name": "America/Bahia"})
    if p == f"act_{acc}/insights":
        if not STATE["rate"]: STATE["rate"] = True; return err(400, 17, "User request limit reached")
        tr = json.loads(q["time_range"][0]); since, until = date.fromisoformat(tr["since"]), date.fromisoformat(tr["until"])
        level = q["level"][0]; daily = q["time_increment"][0] == "1"; bks = q["breakdowns"][0].split(",") if q.get("breakdowns") else []
        if daily and bks and (until - since).days >= 4:
            STATE["too_much"] += 1; return err(500, 1, "Please reduce the amount of data you're asking for, then retry your request")
        statuses = None
        for f in json.loads(q.get("filtering", ["[]"])[0]):
            if f["field"].endswith(".effective_status"): statuses = set(f["value"])
        facts = [f for f in FACTS[cli] if since <= f["day"] <= until]
        if level == "ad": facts = [f for f in facts if (f["ad"][7] in statuses if statuses else f["ad"][7] != "ARCHIVED")]
        rows = agg(facts, level, daily, bks, since, until, acc)
        if "video_play_actions" in q["fields"][0]:
            for r in rows:  # vídeo: plays = impressões/4, cada quartil um pouco menor (lista de ações video_view)
                base = int(r["impressions"]) // 4
                for i, campo in enumerate(("video_play_actions", "video_p25_watched_actions", "video_p50_watched_actions",
                                           "video_p75_watched_actions", "video_p100_watched_actions", "video_thruplay_watched_actions")):
                    r[campo] = [{"action_type": "video_view", "value": str(base - i * 10)}]
        limit = int(q.get("limit", ["25"])[0]); off = int(q.get("after", ["0"])[0])
        body = {"data": rows[off:off + limit], "paging": {"cursors": {"after": str(off + limit)}}}
        if off + limit < len(rows): body["paging"]["next"] = "https://graph.facebook.com/x?access_token=NAO_SEGUIR"
        return httpx.Response(200, json=body)
    if p == f"act_{acc}/campaigns":
        return httpx.Response(200, json={"data": [{"id": a[0], "name": a[1], "objective": a[2], "effective_status": "ACTIVE", "daily_budget": "50000"} for a in ADS[cli]]})
    if p == f"act_{acc}/adsets":
        return httpx.Response(200, json={"data": [{"id": a[3], "name": a[4], "campaign_id": a[0], "effective_status": "ACTIVE", "bid_amount": "600", "bid_strategy": "COST_CAP"} for a in ADS[cli]]})
    if p == f"act_{acc}/ads":
        assert "object_type" in q["fields"][0] and "video_id" in q["fields"][0], q["fields"][0]
        return httpx.Response(200, json={"data": [{"id": a[5], "name": a[6], "adset_id": a[3], "campaign_id": a[0], "effective_status": a[7],
                                                   "creative": {"id": a[5] + "9", "object_type": "VIDEO" if i % 2 else "SHARE",
                                                                **({"video_id": "77" + a[5][-3:]} if i % 2 else {})}}
                                                  for i, a in enumerate(ADS[cli], 1)]})
    return err(400, 100, f"Unsupported get request on {p}")

server._TRANSPORT = httpx.MockTransport(handler)
T = {c: dict(spend=sum(f["spend"] for f in FACTS[c]), buys=sum(f["buys"] for f in FACTS[c]), reach=len(set().union(*(f["people"] for f in FACTS[c])))) for c in FACTS}
ok = []
def check(name, cond, detail=""):
    ok.append(bool(cond)); print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))
def vazou(txt): return any(t in txt for t in TOK.values())

async def main():
    out = await server.listar_clientes(); print(out)
    check("listar_clientes: 2 clientes, sem tokens", "2 cliente(s)" in out and "primos,Primos Premiações" in out and "donaco" in out and not vazou(out))
    out = await server.verificar_conexao(); print(out, "\n")
    check("verificar_conexao sem cliente: testa todos", "[primos] Conexão OK. Token de: sysuser-primos" in out and "[donaco] Conexão OK. Token de: sysuser-donaco" in out)

    out = await server.consultar_insights(desde="2026-09-01", ate="2026-09-14")
    check("2 clientes e nenhum informado: recusa em vez de escolher", out.startswith("ERRO") and "Informe qual" in out and "primos, donaco" in out.replace("donaco, primos", "primos, donaco"), out[:120])

    n0 = len(LOG)
    out = await server.consultar_insights(cliente="primos", desde="2026-09-01", ate="2026-09-14", nivel="campanha"); print(out[:330], "\n")
    check("primos: usa a conta padrão dele e os números dele", f"Cliente: primos | Conta: PRIMOS BM01 CA 01 ({ACC['primos']})" in out and f"gasto BRL {T['primos']['spend']:,.2f}" in out and f"compras {T['primos']['buys']:,.0f}" in out)
    check("primos: todas as chamadas foram com o token do primos", {c for c, _ in LOG[n0:]} == {"primos"})
    check("limite de chamadas: tentou de novo", STATE["rate"])

    n0 = len(LOG)
    out = await server.consultar_insights(cliente="Dona Co", conta="donaco", desde="2026-09-01", ate="2026-09-14", nivel="campanha")
    check("donaco (pelo nome de exibição 'Dona Co'): números dele", "Cliente: donaco" in out and f"gasto BRL {T['donaco']['spend']:,.2f}" in out, out.splitlines()[0])
    check("donaco: todas as chamadas foram com o token do donaco", {c for c, _ in LOG[n0:]} == {"donaco"})
    out = await server.consultar_insights(cliente="donaco", desde="2026-09-01", ate="2026-09-14")
    check("donaco sem conta padrão: pede a conta, não usa a do primos", out.startswith("ERRO") and "não tem conta padrão" in out)
    out = await server.consultar_insights(cliente="donaco", conta=ACC["primos"], desde="2026-09-01", ate="2026-09-14")
    check("token do donaco não lê a conta do primos", out.startswith("ERRO") and "permissão" in out and f"{T['primos']['spend']:,.2f}" not in out, out[:90])
    out = await server.consultar_insights(cliente="donaco", conta="primos")
    check("donaco não acha conta 'primos' pelo nome", out.startswith("ERRO") and "Nenhuma conta com" in out)

    # chamadas simultâneas de clientes diferentes não se cruzam
    n0 = len(LOG)
    a, b = await asyncio.gather(
        server.consultar_insights(cliente="primos", desde="2026-09-01", ate="2026-09-14", nivel="anuncio", por_dia=True, quebra="idade_genero", max_linhas=3),
        server.consultar_insights(cliente="donaco", conta=ACC["donaco"], desde="2026-09-01", ate="2026-09-14", nivel="anuncio", por_dia=True, quebra="idade_genero", max_linhas=3))
    cruzou = [(c, p) for c, p in LOG[n0:] if p.startswith("act_") and not p.startswith(f"act_{ACC[c]}")]
    check("simultâneas: cada token só tocou a própria conta", not cruzou and f"{T['primos']['spend']:,.2f}" in a and f"{T['donaco']['spend']:,.2f}" in b, f"{len(LOG) - n0} requisições, cruzadas={len(cruzou)}")
    check("simultâneas: período dividido e paginado", STATE["too_much"] > 0 and "Linhas: 252" in a and "Linhas: 168" in b, f"{a.splitlines()[1][:12]} / {b.splitlines()[1][:12]}")

    out = await server.consultar_insights(cliente="primos", desde="2026-09-01", ate="2026-09-14", nivel="anuncio", por_dia=True, com_video=True, max_linhas=3000)
    linhas = out.split("\n\n", 1)[1].splitlines()
    cab = linhas[0].split(",")
    check("com_video: colunas de vídeo no fim, na ordem", cab[-6:] == ["video_plays", "video_p25", "video_p50", "video_p75", "video_p100", "thruplays"], ",".join(cab[-6:]))
    import csv as _csv, io as _io
    rr = list(_csv.DictReader(_io.StringIO(out.split("\n\n", 1)[1])))
    check("com_video: valores inteiros e coerentes (plays = impressões/4, quartis caindo)",
          all(int(r["video_plays"]) == int(r["impressoes"]) // 4 and int(r["video_p25"]) == int(r["video_plays"]) - 10
              and int(r["thruplays"]) == int(r["video_plays"]) - 50 for r in rr) and len(rr) == 42, f"{len(rr)} linhas")
    out = await server.consultar_insights(cliente="primos", desde="2026-09-01", ate="2026-09-14", nivel="anuncio", por_dia=True, max_linhas=3)
    check("sem com_video: nada muda (sem colunas de vídeo)", "video_plays" not in out and out.split("\n\n", 1)[1].splitlines()[0].endswith(",ticket"))
    out = await server.estrutura_conta(cliente="primos", nivel="anuncios", somente_ativos=False)
    est = json.loads(next(l for l in out.splitlines() if l.startswith("[")))
    check("estrutura anúncios: tipo do criativo e video_id", [e.get("tipo_criativo") for e in est] == ["VIDEO", "SHARE", "VIDEO"]
          and est[0].get("video_id") and "video_id" not in est[1], json.dumps(est[:2], ensure_ascii=False)[:160])

    out = await server.estrutura_conta(cliente="primos", nivel="conjuntos")
    check("estrutura: limite de CPA em reais", "Cliente: primos" in out and '"limite_cpa_ou_lance":6.0' in out)

    outs = {}
    for cli in ("primos", "donaco"):
        outs[cli] = await server.exportar_relatorio(cliente=cli, conta=ACC[cli], desde="2026-09-01", ate="2026-09-14")
    print(outs["primos"], "\n")
    from openpyxl import load_workbook
    for cli in ("primos", "donaco"):
        path = Path(outs[cli].splitlines()[0].split(": ", 1)[1])
        check(f"export {cli}: arquivo na pasta do cliente", path.parent == tmp / "relatorios" / cli and path.name.startswith(f"meta-ads_{cli}_"), str(path.relative_to(tmp)))
        wb = load_workbook(path, read_only=True)
        a = list(wb["Tabela A - Idade e Genero"].iter_rows(values_only=True)); h = list(a[0]); gi, ri = h.index("Valor usado (BRL)"), h.index("Alcance")
        check(f"export {cli}: soma das linhas = total do próprio cliente (com arquivados)", abs(sum(r[gi] for r in a[2:]) - T[cli]["spend"]) < 0.05 and a[1][ri] == T[cli]["reach"])
        info = dict((r[0], r[1]) for r in wb["Info"].iter_rows(values_only=True))
        check(f"export {cli}: aba Info identifica o cliente", info.get("Cliente") == cli)
        ids = {r[h.index("Identificação da campanha")] for r in a[2:]}
        check(f"export {cli}: só campanhas dele", ids == {x[0] for x in ADS[cli]})
    check("export: conferência 0,00%", all("gasto A +0.00%" in o for o in outs.values()))

    (tmp / "clientes/ruim.env").write_text(f"META_ACCESS_TOKEN={TOK['ruim']}\n")
    out = await server.listar_contas(cliente="ruim"); print(out)
    check("token inválido: erro cita o cliente e não vaza token", out.startswith("ERRO") and "cliente 'ruim'" in out and not vazou(out))
    out = await server.verificar_conexao()
    check("um cliente com token ruim não derruba os outros", "[primos] Conexão OK" in out and "[ruim] ERRO" in out and not vazou(out))
    out = await server.consultar_insights(cliente="inexistente")
    check("cliente inexistente: lista os configurados", out.startswith("ERRO") and "não está configurado" in out and "primos" in out)
    check("nenhuma saída vazou token", not any(vazou(o) for o in list(outs.values()) + [a, b]))
    check("somente GET, token nunca em URL (asserts no servidor falso)", True)

    for f in (tmp / "clientes").glob("*.env"): f.unlink()
    out = await server.listar_contas()
    check("sem clientes: orienta a rodar o configurador", out.startswith("ERRO") and "Nenhum cliente configurado" in out and "configurar.py" in out)
    (tmp / ".env").write_text("META_ACCESS_TOKEN=" + TOK["primos"] + "\nMETA_AD_ACCOUNT_ID=" + ACC["primos"] + "\n")
    out = await server.consultar_insights(desde="2026-09-01", ate="2026-09-14")
    check("instalação antiga (.env único) segue funcionando como 'padrao'", "Cliente: padrao" in out and f"{T['primos']['spend']:,.2f}" in out)

asyncio.run(main())
print(f"\n{sum(ok)}/{len(ok)} verificações passaram"); sys.exit(0 if all(ok) else 1)
